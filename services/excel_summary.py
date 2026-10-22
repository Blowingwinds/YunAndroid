"""Configuration and execution service for the resident Excel summary job."""

from __future__ import annotations

import configparser
import json
import re
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Callable
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from services.feishu_notifier import FeishuNotifier
from services.monitoring import BEIJING_TZ


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG_PATH = PROJECT_ROOT / "config" / "monitor.ini"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "data" / "excel" / "Bydate"
DEFAULT_STATE_FILE = PROJECT_ROOT / "data" / "monitor" / "excel_summary_state.json"

SUMMARY_HEADERS = ("设备ID", "应用名", "设备区组", "包名", "状态", "说明", "来源文件")
REPORT_HEADERS = (
    "日期",
    "应用名",
    "设备区组",
    "包名",
    "JSON内区组",
    "类型",
    "角色GID",
    "角色名",
    "等级",
    "家族",
    "消息",
    "线路",
    "频道",
    "时间",
    "@引用",
    "图标",
    "来源文件",
)


@dataclass(frozen=True)
class ExcelSummaryConfig:
    config_path: Path
    queue_file: Path
    json_root: Path
    output_dir: Path
    state_file: Path
    run_time: str
    source_date: str
    tag: str
    max_workers: int
    catchup_seconds: int
    device_lock_timeout_seconds: int
    control_host: str
    control_port: int
    control_token: str
    feishu_webhook_url: str
    feishu_security_keyword: str
    feishu_secret: str

    @property
    def run_clock(self) -> tuple[int, int]:
        return parse_clock(self.run_time)


def _positive_int(parser: configparser.ConfigParser, section: str, key: str, fallback: int) -> int:
    try:
        value = parser.getint(section, key, fallback=fallback)
    except ValueError as exc:
        raise ValueError(f"配置项 [{section}] {key} 必须是整数") from exc
    if value <= 0:
        raise ValueError(f"配置项 [{section}] {key} 必须大于 0")
    return value


def _resolve_path(config_dir: Path, value: str, fallback: Path) -> Path:
    text = str(value).strip()
    path = Path(text) if text else fallback
    if not path.is_absolute():
        path = config_dir / path
    return path.resolve()


def parse_clock(value: str) -> tuple[int, int]:
    text = str(value).strip()
    try:
        parsed = datetime.strptime(text, "%H:%M")
    except ValueError as exc:
        raise ValueError("配置项 [excel_summary] run_time 必须是 HH:MM") from exc
    return parsed.hour, parsed.minute


def load_excel_summary_config(path: str | Path = DEFAULT_CONFIG_PATH) -> ExcelSummaryConfig:
    config_path = Path(path).resolve()
    if not config_path.is_file():
        raise FileNotFoundError(f"汇总配置文件不存在：{config_path}")

    parser = configparser.ConfigParser(interpolation=None)
    try:
        with config_path.open("r", encoding="utf-8-sig") as handle:
            parser.read_file(handle)
    except (OSError, ValueError) as exc:
        raise ValueError(f"读取汇总配置失败：{exc}") from exc

    run_time = parser.get("excel_summary", "run_time", fallback="09:40").strip()
    parse_clock(run_time)
    source_date = parser.get("excel_summary", "source_date", fallback="yesterday").strip()
    resolve_source_date(source_date, datetime.now(BEIJING_TZ))
    tag = parser.get("excel_summary", "tag", fallback="日常监控").strip()
    if not tag:
        raise ValueError("配置项 [excel_summary] tag 不能为空")
    max_workers = _positive_int(parser, "excel_summary", "max_workers", 8)
    if max_workers > 16:
        raise ValueError("配置项 [excel_summary] max_workers 不能超过 16")

    catchup_minutes = _positive_int(
        parser, "excel_summary", "catchup_minutes", 120
    )
    lock_timeout_minutes = _positive_int(
        parser, "excel_summary", "device_lock_timeout_minutes", 30
    )
    control_port = _positive_int(
        parser, "excel_summary_control", "control_port", 18766
    )
    if control_port > 65535:
        raise ValueError("配置项 [excel_summary_control] control_port 不能超过 65535")
    control_host = parser.get(
        "excel_summary_control", "control_host", fallback="127.0.0.1"
    ).strip()
    if control_host != "127.0.0.1":
        raise ValueError("汇总控制接口只能绑定 127.0.0.1")
    feishu_security_keyword = parser.get(
        "feishu", "security_keyword", fallback=""
    ).strip()
    if not feishu_security_keyword:
        raise ValueError("配置项 [feishu] security_keyword 不能为空")

    config_dir = config_path.parent
    return ExcelSummaryConfig(
        config_path=config_path,
        queue_file=_resolve_path(
            config_dir,
            parser.get("excel_summary_paths", "queue_file", fallback="cloud_phone_queue.xlsx"),
            config_dir / "cloud_phone_queue.xlsx",
        ),
        json_root=_resolve_path(
            config_dir,
            parser.get("excel_summary_paths", "json_root", fallback="../SF_json"),
            PROJECT_ROOT / "SF_json",
        ),
        output_dir=_resolve_path(
            config_dir,
            parser.get("excel_summary_paths", "output_dir", fallback="../data/excel/Bydate"),
            DEFAULT_OUTPUT_DIR,
        ),
        state_file=_resolve_path(
            config_dir,
            parser.get(
                "excel_summary_paths",
                "state_file",
                fallback="../data/monitor/excel_summary_state.json",
            ),
            DEFAULT_STATE_FILE,
        ),
        run_time=run_time,
        source_date=source_date,
        tag=tag,
        max_workers=max_workers,
        catchup_seconds=catchup_minutes * 60,
        device_lock_timeout_seconds=lock_timeout_minutes * 60,
        control_host=control_host,
        control_port=control_port,
        control_token=parser.get(
            "excel_summary_control", "control_token", fallback=""
        ).strip(),
        feishu_webhook_url=parser.get(
            "feishu",
            "alert_webhook_url",
            fallback=parser.get("feishu", "webhook_url", fallback=""),
        ).strip(),
        feishu_security_keyword=feishu_security_keyword,
        feishu_secret=parser.get("feishu", "secret", fallback="").strip(),
    )


def resolve_source_date(value: str, now: datetime | None = None) -> str:
    moment = now or datetime.now(BEIJING_TZ)
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=BEIJING_TZ)
    else:
        moment = moment.astimezone(BEIJING_TZ)
    text = str(value).strip().lower()
    if text in {"today", "今天"}:
        return moment.date().isoformat()
    if text in {"yesterday", "昨天"}:
        return (moment.date() - timedelta(days=1)).isoformat()
    candidate = text.replace("/", "-").replace(".", "-")
    for pattern in ("%Y-%m-%d", "%Y%m%d"):
        try:
            return datetime.strptime(candidate, pattern).date().isoformat()
        except ValueError:
            continue
    raise ValueError(
        f"配置项 [excel_summary] source_date 格式不正确：{value}"
    )


def build_summary_message(summary: dict[str, object], security_keyword: str) -> str:
    return "\n".join(
        (
            f"From : {security_keyword}",
            "【云手机每日汇总】",
            f"数据日期：{summary.get('source_date', '')}",
            f"完成时间：{summary.get('finished_at', '')}",
            (
                "设备：总计 {total}，成功 {success}，失败 {failed}，跳过 {skipped}"
            ).format(
                total=summary.get("total", 0),
                success=summary.get("success", 0),
                failed=summary.get("failed", 0),
                skipped=summary.get("skipped", 0),
            ),
            f"汇总记录：{summary.get('row_count', 0)} 行",
            f"文件：{summary.get('output_file', '')}",
        )
    )


def parse_daily_json(path: str | Path, target_date: str) -> list[dict[str, object]]:
    source = Path(path)
    text = source.read_text(encoding="utf-8")
    text = re.sub(r"[\x00-\x08\x0B-\x1F\x7F]", "", text.replace(r"\'", "'"))
    payload = json.loads(text)
    if not isinstance(payload, dict):
        raise ValueError(f"JSON 顶层必须是对象：{source}")

    rows: list[dict[str, object]] = []
    for zone_name, zone_data in payload.items():
        if not isinstance(zone_data, dict):
            continue
        chars = zone_data.get("chars")
        if not isinstance(chars, dict):
            continue
        date_content = chars.get(target_date)
        if not isinstance(date_content, dict):
            continue

        for gid, char_info in date_content.items():
            if gid == "messages" or not isinstance(char_info, dict):
                continue
            message = char_info.get("msg", "")
            rows.append(
                {
                    "日期": target_date,
                    "JSON内区组": str(zone_name),
                    "类型": "人物发言" if str(message).strip() else "角色信息",
                    "角色GID": str(gid),
                    "角色名": char_info.get("name", ""),
                    "等级": char_info.get("level", ""),
                    "家族": char_info.get("family", ""),
                    "消息": message,
                    "线路": "",
                    "频道": "",
                    "时间": "",
                    "@引用": "",
                    "图标": char_info.get("icon", ""),
                    "来源文件": source.name,
                }
            )

        messages = date_content.get("messages", [])
        if not isinstance(messages, list):
            continue
        for message in messages:
            if not isinstance(message, dict):
                continue
            timestamp = message.get("time")
            time_text = ""
            if timestamp:
                try:
                    time_text = datetime.fromtimestamp(
                        int(timestamp), BEIJING_TZ
                    ).strftime("%Y-%m-%d %H:%M:%S")
                except (TypeError, ValueError, OSError):
                    time_text = str(timestamp)
            card_list = message.get("cardList", [])
            rows.append(
                {
                    "日期": target_date,
                    "JSON内区组": str(zone_name),
                    "类型": "聊天消息",
                    "角色GID": message.get("gid", ""),
                    "角色名": message.get("name", ""),
                    "等级": message.get("level", ""),
                    "家族": "",
                    "消息": message.get("msg", ""),
                    "线路": message.get("server_name", ""),
                    "频道": message.get("channel", ""),
                    "时间": time_text,
                    "@引用": " | ".join(str(item) for item in card_list)
                    if isinstance(card_list, list)
                    else str(card_list or ""),
                    "图标": message.get("icon", ""),
                    "来源文件": source.name,
                }
            )
    return rows


def _format_worksheet(worksheet, headers: tuple[str, ...]) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for index, header in enumerate(headers, start=1):
        width = 42 if header == "消息" else max(12, min(28, len(header) * 2 + 4))
        worksheet.column_dimensions[get_column_letter(index)].width = width


def write_summary_workbook(
    pull_summary: dict,
    source_date: str,
    output_file: str | Path,
) -> dict[str, object]:
    results = list(pull_summary.get("results") or [])
    if not results:
        raise RuntimeError("没有匹配的设备，不生成汇总文件")
    successful_sources = sum(
        1 for item in results if str(item.get("status")) == "成功"
    )
    if not successful_sources:
        raise RuntimeError("所有设备均拉取失败，保留已有汇总文件")

    destination = Path(output_file)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "_执行摘要"
    summary_sheet.append(SUMMARY_HEADERS)
    report_sheet = workbook.create_sheet(source_date)
    report_sheet.append(REPORT_HEADERS)

    row_count = 0
    parse_failures = 0
    for result in results:
        source_path = Path(str(result.get("local_path") or ""))
        status = str(result.get("status") or "")
        detail = str(result.get("detail") or "")
        parsed_rows: list[dict[str, object]] = []
        if status == "成功":
            try:
                parsed_rows = parse_daily_json(source_path, source_date)
            except (OSError, ValueError, json.JSONDecodeError) as exc:
                status = "解析失败"
                detail = str(exc)
                parse_failures += 1

        summary_sheet.append(
            (
                result.get("device_id", ""),
                result.get("app_name", ""),
                result.get("group", ""),
                result.get("package_name", ""),
                status,
                detail,
                str(source_path) if str(source_path) != "." else "",
            )
        )
        for row in parsed_rows:
            values = {
                **row,
                "应用名": result.get("app_name", ""),
                "设备区组": result.get("group", ""),
                "包名": result.get("package_name", ""),
            }
            report_sheet.append(tuple(values.get(header, "") for header in REPORT_HEADERS))
            row_count += 1

    if parse_failures == successful_sources:
        workbook.close()
        raise RuntimeError("所有成功拉取的 JSON 均解析失败，保留已有汇总文件")

    _format_worksheet(summary_sheet, SUMMARY_HEADERS)
    _format_worksheet(report_sheet, REPORT_HEADERS)

    temporary = destination.with_name(
        f".{destination.stem}.{uuid4().hex}.tmp.xlsx"
    )
    try:
        workbook.save(temporary)
        temporary.replace(destination)
    finally:
        workbook.close()
        if temporary.exists():
            temporary.unlink()
    return {
        "output_file": str(destination),
        "row_count": row_count,
        "parse_failures": parse_failures,
    }


class ExcelSummaryRunner:
    def __init__(
        self,
        config: ExcelSummaryConfig,
        *,
        daily_runner_factory: Callable[[ExcelSummaryConfig], object] | None = None,
        notifier: FeishuNotifier | None = None,
    ) -> None:
        self.config = config
        self.daily_runner_factory = daily_runner_factory or self._build_daily_runner
        self.notifier = notifier or FeishuNotifier(
            config.feishu_webhook_url,
            secret=config.feishu_secret,
        )

    @staticmethod
    def _build_daily_runner(config: ExcelSummaryConfig):
        from tools.daily_monitor import DailyMonitorRunner

        return DailyMonitorRunner(
            queue_file=config.queue_file,
            output_root=config.json_root,
            tag=config.tag,
            max_workers=config.max_workers,
        )

    def run_once(self, now: datetime | None = None) -> dict[str, object]:
        started = now or datetime.now(BEIJING_TZ)
        if started.tzinfo is None:
            started = started.replace(tzinfo=BEIJING_TZ)
        else:
            started = started.astimezone(BEIJING_TZ)
        source_date = resolve_source_date(self.config.source_date, started)
        pull_summary = self.daily_runner_factory(self.config).run(source_date)
        output_file = self.config.output_dir / f"{source_date}.xlsx"
        workbook_summary = write_summary_workbook(pull_summary, source_date, output_file)
        finished = datetime.now(BEIJING_TZ)
        summary: dict[str, object] = {
            "run_date": started.date().isoformat(),
            "source_date": source_date,
            "tag": self.config.tag,
            "started_at": started.isoformat(timespec="seconds"),
            "finished_at": finished.isoformat(timespec="seconds"),
            "output_file": workbook_summary["output_file"],
            "row_count": workbook_summary["row_count"],
            "parse_failures": workbook_summary["parse_failures"],
            "total": int(pull_summary.get("total", 0) or 0),
            "success": int(pull_summary.get("success", 0) or 0),
            "failed": int(pull_summary.get("failed", 0) or 0),
            "skipped": int(pull_summary.get("skipped", 0) or 0),
            "notification_error": "",
        }
        try:
            self.notifier.send_text(
                build_summary_message(summary, self.config.feishu_security_keyword)
            )
        except Exception as exc:  # noqa: BLE001 - 汇总文件成功不因通知失败而回滚
            summary["notification_error"] = str(exc)
        self._write_state(summary)
        return summary

    def _write_state(self, summary: dict[str, object]) -> None:
        path = self.config.state_file
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary = path.with_suffix(path.suffix + ".tmp")
        temporary.write_text(
            json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        temporary.replace(path)
