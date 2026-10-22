"""Persistent cloud-phone connection queue storage."""

from __future__ import annotations

import csv
import os
from dataclasses import dataclass, replace
from ipaddress import IPv4Address
from pathlib import Path

from openpyxl import Workbook, load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
QUEUE_FILE = PROJECT_ROOT / "config" / "cloud_phone_queue.xlsx"
_PUBLIC_IP_VALUES = [value.strip() for value in os.environ.get("YUNANDROID_PUBLIC_IPS", "127.0.0.1,127.0.0.2").split(",")]
PUBLIC_IP_PRESETS = tuple((_PUBLIC_IP_VALUES + ["127.0.0.2"])[:2])
DEFAULT_PUBLIC_IP = PUBLIC_IP_PRESETS[0]
QUEUE_FIELDNAMES = [
    "id",
    "enabled",
    "sort_order",
    "app_name",
    "public_ip",
    "remote_ip",
    "remote_port",
    "local_port",
    "note",
    "last_connected_at",
    "connect_count",
    "package_name",
    "tag",
    "group",
]
REQUIRED_COLUMNS = {
    "id",
    "enabled",
    "sort_order",
    "app_name",
    "remote_ip",
    "remote_port",
    "local_port",
    "package_name",
}


@dataclass(frozen=True)
class Device:
    id: str
    enabled: bool
    sort_order: int
    app_name: str
    remote_ip: str
    remote_port: int
    local_port: int
    note: str = ""
    last_connected_at: str = ""
    connect_count: int = 0
    package_name: str = ""
    tag: str = ""
    group: str = ""
    public_ip: str = DEFAULT_PUBLIC_IP

    @property
    def serial(self) -> str:
        return f"127.0.0.1:{self.local_port}"

    @property
    def remote_address(self) -> str:
        return f"{self.remote_ip}:{self.remote_port}"

    @property
    def label(self) -> str:
        return self.app_name or self.remote_address


def _cell_text(value: object) -> str:
    """Normalize an Excel/CSV cell for text fields without turning blanks into 'None'."""
    return "" if value is None else str(value).strip()


def normalize_public_ip(value: object, *, row_number: int | None = None) -> str:
    """Return a canonical IPv4 address for one public SSH server."""
    text = _cell_text(value)
    try:
        return str(IPv4Address(text))
    except ValueError as exc:
        location = f"row {row_number}" if row_number is not None else "connection record"
        raise ValueError(
            f"{location} public_ip must be a valid IPv4 address: {text or '<blank>'}"
        ) from exc


def public_ip_options(devices: list[Device]) -> list[str]:
    """Return configured and bootstrap public IPs in numeric address order."""
    values = {
        normalize_public_ip(value)
        for value in (
            *(device.public_ip for device in devices),
            *PUBLIC_IP_PRESETS,
        )
    }
    return sorted(values, key=IPv4Address)


def _parse_bool(value: object, field: str, row_number: int) -> bool:
    if isinstance(value, bool):
        return value
    normalized = _cell_text(value).lower()
    if normalized in {"1", "true", "yes", "y"}:
        return True
    if normalized in {"0", "false", "no", "n"}:
        return False
    raise ValueError(f"第 {row_number} 行字段 {field} 必须是 0 或 1")


def _parse_int(value: object, field: str, row_number: int) -> int:
    try:
        if isinstance(value, bool):
            raise ValueError
        if isinstance(value, float) and not value.is_integer():
            raise ValueError
        return int(value) if isinstance(value, (int, float)) else int(_cell_text(value))
    except ValueError as exc:
        raise ValueError(f"第 {row_number} 行字段 {field} 必须是整数") from exc


def load_devices(path: str | Path = QUEUE_FILE) -> list[Device]:
    """Load the persistent connection queue in queue order."""
    queue_path = Path(path)
    if not queue_path.exists():
        raise FileNotFoundError(f"连接队列不存在：{queue_path}")

    if queue_path.suffix.lower() == ".csv":
        rows = _read_csv_rows(queue_path)
    else:
        rows = _read_xlsx_rows(queue_path)

    devices = []
    for row_number, row in rows:
        if not any(_cell_text(value) for value in row.values()):
            continue
        public_ip = (
            DEFAULT_PUBLIC_IP
            if "public_ip" not in row
            else normalize_public_ip(row.get("public_ip"), row_number=row_number)
        )
        devices.append(
            Device(
                id=_cell_text(row.get("id")),
                enabled=_parse_bool(row.get("enabled"), "enabled", row_number),
                sort_order=_parse_int(row.get("sort_order"), "sort_order", row_number),
                app_name=_cell_text(row.get("app_name")),
                remote_ip=_cell_text(row.get("remote_ip")),
                remote_port=_parse_int(row.get("remote_port"), "remote_port", row_number),
                local_port=_parse_int(row.get("local_port"), "local_port", row_number),
                note=_cell_text(row.get("note")),
                last_connected_at=_cell_text(row.get("last_connected_at")),
                connect_count=_parse_int(row.get("connect_count") or 0, "connect_count", row_number),
                package_name=_cell_text(row.get("package_name")),
                tag=_cell_text(row.get("tag")),
                group=_cell_text(row.get("group")),
                public_ip=public_ip,
            )
        )
    return sorted(devices, key=lambda device: (device.sort_order, device.id))


def save_devices(path: str | Path, devices: list[Device]) -> None:
    """Write the queue atomically as Excel, with legacy CSV support for explicit .csv paths."""
    queue_path = Path(path)
    queue_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = queue_path.with_suffix(queue_path.suffix + ".tmp")
    try:
        if queue_path.suffix.lower() == ".csv":
            _write_csv(temporary_path, devices)
        else:
            _write_xlsx(temporary_path, devices)
        temporary_path.replace(queue_path)
    except Exception:
        temporary_path.unlink(missing_ok=True)
        raise


def _validate_columns(columns: list[object]) -> None:
    normalized = {_cell_text(column) for column in columns if _cell_text(column)}
    missing = REQUIRED_COLUMNS - normalized
    if missing:
        raise ValueError(f"连接队列缺少字段：{', '.join(sorted(missing))}")


def _read_csv_rows(path: Path) -> list[tuple[int, dict[str, object]]]:
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        _validate_columns(list(reader.fieldnames or []))
        return [(row_number, dict(row)) for row_number, row in enumerate(reader, 2)]


def _read_xlsx_rows(path: Path) -> list[tuple[int, dict[str, object]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        row_iterator = worksheet.iter_rows(values_only=True)
        try:
            headers = list(next(row_iterator))
        except StopIteration:
            _validate_columns([])
            return []
        _validate_columns(headers)
        normalized_headers = [_cell_text(header) for header in headers]
        rows = []
        for row_number, values in enumerate(row_iterator, 2):
            rows.append((row_number, dict(zip(normalized_headers, values))))
        return rows
    finally:
        workbook.close()


def _device_row(device: Device) -> dict[str, object]:
    return {
        "id": device.id,
        "enabled": int(device.enabled),
        "sort_order": device.sort_order,
        "app_name": device.app_name,
        "public_ip": normalize_public_ip(device.public_ip),
        "remote_ip": device.remote_ip,
        "remote_port": device.remote_port,
        "local_port": device.local_port,
        "note": device.note,
        "last_connected_at": device.last_connected_at,
        "connect_count": device.connect_count,
        "package_name": device.package_name,
        "tag": device.tag,
        "group": device.group,
    }


def _write_csv(path: Path, devices: list[Device]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=QUEUE_FIELDNAMES)
        writer.writeheader()
        writer.writerows(_device_row(device) for device in devices)


def _write_xlsx(path: Path, devices: list[Device]) -> None:
    workbook = Workbook()
    try:
        worksheet = workbook.active
        worksheet.title = "Devices"
        worksheet.append(QUEUE_FIELDNAMES)
        for device in devices:
            row = _device_row(device)
            worksheet.append([row[field] for field in QUEUE_FIELDNAMES])
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        workbook.save(path)
    finally:
        workbook.close()


def renumber_local_ports(devices: list[Device], start_port: int = 1) -> list[Device]:
    """Return queue records with sequential local ports, preserving all other fields."""

    if not isinstance(start_port, int) or isinstance(start_port, bool):
        raise ValueError("本地起始端口必须是整数")
    if not 1 <= start_port <= 65535:
        raise ValueError(f"本地起始端口超出范围: {start_port}")
    records = list(devices)
    end_port = start_port + len(records) - 1
    if end_port > 65535:
        raise ValueError("本地端口范围超出 65535")
    return [replace(device, local_port=start_port + index) for index, device in enumerate(records)]


def next_device_id(devices: list[Device]) -> str:
    """Return the next readable queue ID without reusing an existing ID."""
    used = {device.id for device in devices}
    number = 1
    while f"cp-{number:04d}" in used:
        number += 1
    return f"cp-{number:04d}"
